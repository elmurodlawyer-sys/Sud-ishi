"""Jadvallarni Excel va PDF formatlarida eksport qilish hamda chop etish (3.8-band)."""
import io
from datetime import date, datetime
from decimal import Decimal

from django.conf import settings
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone
from django.utils.text import slugify

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_fonts_registered = False


def _register_fonts():
    global _fonts_registered
    if _fonts_registered:
        return
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    font_dir = settings.BASE_DIR / "static" / "fonts"
    pdfmetrics.registerFont(TTFont("DejaVu", str(font_dir / "DejaVuSans.ttf")))
    pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(font_dir / "DejaVuSans-Bold.ttf")))
    pdfmetrics.registerFontFamily("DejaVu", normal="DejaVu", bold="DejaVu-Bold")
    _fonts_registered = True


def fmt_value(value):
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return f"{value:,.2f}".replace(",", " ")
    if isinstance(value, float):
        return f"{value:,.2f}".replace(",", " ")
    if isinstance(value, int) and not isinstance(value, bool):
        return f"{value:,}".replace(",", " ")
    if isinstance(value, datetime):
        return timezone.localtime(value).strftime("%d.%m.%Y %H:%M") if timezone.is_aware(value) else value.strftime("%d.%m.%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    return str(value)


def _excel_value(value):
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime) and timezone.is_aware(value):
        return timezone.localtime(value).replace(tzinfo=None)
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool, date)):
        return value
    return str(value)


def build_xlsx(title, headers, rows, filters=None, totals=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Hisobot"[:31]
    ncols = max(len(headers), 1)
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.cell(row=2, column=1, value=f"Shakllantirilgan: {timezone.localtime():%d.%m.%Y %H:%M}").font = Font(italic=True, size=9)
    row_idx = 3
    for label, value in filters or []:
        ws.cell(row=row_idx, column=1, value=f"{label}: {value}").font = Font(size=9)
        row_idx += 1
    row_idx += 1
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="1F3A5F")
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row_idx, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = border
    header_row = row_idx
    for row in rows:
        row_idx += 1
        for col, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col, value=_excel_value(value))
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(value, datetime):
                cell.number_format = "DD.MM.YYYY HH:MM"
            elif isinstance(value, date):
                cell.number_format = "DD.MM.YYYY"
            elif isinstance(value, (Decimal, float)):
                cell.number_format = "#,##0.00"
    if totals:
        row_idx += 1
        for col, value in enumerate(totals, start=1):
            cell = ws.cell(row=row_idx, column=col, value=_excel_value(value))
            cell.font = Font(bold=True)
            cell.border = border
            if isinstance(value, (Decimal, float)):
                cell.number_format = "#,##0.00"
    for col, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(fmt_value(r[col - 1])) for r in rows[:500] if col - 1 < len(r)] + [8])
        ws.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 10), 60)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_pdf(title, headers, rows, filters=None, totals=None):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    _register_fonts()
    buffer = io.BytesIO()
    page = landscape(A4)
    doc = SimpleDocTemplate(
        buffer, pagesize=page, leftMargin=10 * mm, rightMargin=10 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
        title=title, author="Sud monitoringi axborot tizimi",
    )
    title_style = ParagraphStyle("t", fontName="DejaVu-Bold", fontSize=13, leading=16, spaceAfter=4)
    small = ParagraphStyle("s", fontName="DejaVu", fontSize=8, leading=10)
    cell = ParagraphStyle("c", fontName="DejaVu", fontSize=7.5, leading=9)
    head = ParagraphStyle("h", fontName="DejaVu-Bold", fontSize=7.5, leading=9, textColor=colors.white)
    esc = lambda s: str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")  # noqa: E731

    story = [Paragraph(esc(title), title_style)]
    story.append(Paragraph(f"Shakllantirilgan: {timezone.localtime():%d.%m.%Y %H:%M}", small))
    for label, value in filters or []:
        story.append(Paragraph(f"<b>{esc(label)}:</b> {esc(value)}", small))
    story.append(Spacer(1, 4 * mm))

    data = [[Paragraph(esc(h), head) for h in headers]]
    for row in rows:
        data.append([Paragraph(esc(fmt_value(v)).replace("\n", "<br/>"), cell) for v in row])
    if totals:
        data.append([Paragraph(f"<b>{esc(fmt_value(v))}</b>", cell) for v in totals])
    available = page[0] - 20 * mm
    weights = []
    for i, h in enumerate(headers):
        sample = [len(fmt_value(r[i])) for r in rows[:200] if i < len(r)]
        weights.append(min(max([len(h) * 0.6] + sample + [4]), 45))
    total_w = sum(weights) or 1
    widths = [available * w / total_w for w in weights]
    # Juda tor ustunlar bo'lmasligi uchun minimal kenglik ta'minlanadi
    min_w = min(16 * mm, available / max(len(widths), 1))
    narrow = [w < min_w for w in widths]
    if any(narrow):
        rest_total = sum(w for w, n in zip(widths, narrow) if not n) or 1
        rest_space = available - min_w * sum(narrow)
        widths = [min_w if n else w * rest_space / rest_total for w, n in zip(widths, narrow)]
    table = Table(data, colWidths=widths, repeatRows=1)
    style = [
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5F9")]),
    ]
    if totals:
        style.append(("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E3E9F1")))
    table.setStyle(TableStyle(style))
    story.append(table)

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("DejaVu", 7)
        canvas.drawString(10 * mm, 6 * mm, "“Sud monitoringi” axborot tizimi")
        canvas.drawRightString(page[0] - 10 * mm, 6 * mm, f"{document.page}-bet")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


def export_response(request, fmt, title, headers, rows, filters=None, totals=None, filename="hisobot"):
    """fmt: xlsx | pdf | print"""
    stamp = timezone.localtime().strftime("%Y%m%d_%H%M")
    base = f"{slugify(filename) or 'hisobot'}_{stamp}"
    if fmt == "xlsx":
        response = HttpResponse(build_xlsx(title, headers, rows, filters, totals), content_type=XLSX_TYPE)
        response["Content-Disposition"] = f'attachment; filename="{base}.xlsx"'
    elif fmt == "pdf":
        response = HttpResponse(build_pdf(title, headers, rows, filters, totals), content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="{base}.pdf"'
    else:
        response = render(
            request, "reports/print.html",
            {"title": title, "headers": headers, "rows": [[fmt_value(v) for v in r] for r in rows],
             "totals": [fmt_value(v) for v in totals] if totals else None, "filters": filters, "now": timezone.now()},
        )
    from accounts.audit import log_action

    log_action(request, "eksport", description=f"{title} ({fmt}), {len(rows)} qator")
    return response
