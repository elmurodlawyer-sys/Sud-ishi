"""Tashqi manbalardan sud ishlari ma'lumotlarini olish adapterlari (3.2 va 8-bo'limlar).

Har bir adapter yozuvlarni quyidagi yagona (normallashtirilgan) ko'rinishda qaytaradi:

    {
        "external_key": "...",      # manbadagi yagona kalit (bo'lmasa avtomatik hisoblanadi)
        "case_number": "...",       # sud ishi raqami
        "court_name": "...",        # sud nomi
        "court_type": "...",        # sud turi (matn)
        "region": "...",            # hudud (matn)
        "instance": "...",          # sud instansiyasi (matn)
        "category": "...",          # ish turkumi / turi (matn)
        "plaintiffs": "...",        # da'vogar(lar)
        "defendants": "...",        # javobgar(lar)
        "third_parties": "...",     # uchinchi shaxslar / boshqa ishtirokchilar
        "judge": "...",             # mas'ul sudya
        "hearing_at": "YYYY-MM-DDTHH:MM",  # sud majlisi sanasi va vaqti
        "location": "...",          # majlis zali
        "subject": "...",           # ish predmeti
        "claim_amount": "...",      # da'vo summasi
        "url": "...",               # manbadagi havola
    }

Oliy sud rasmiy API taqdim etganda `JsonApiAdapter` sozlamalar orqali unga
moslashtiriladi. API mavjud bo'lmaganda `HtmlTableAdapter` jadval2.sud.uz
sahifalaridagi jadvallarni o'qiydi.
"""
import csv
import io
import json
import random
import re
from datetime import date, datetime, timedelta
from html.parser import HTMLParser
from urllib.parse import urljoin

import requests
from django.conf import settings
from django.utils import timezone

from core.normalize import normalize_name

FIELDS = [
    "external_key", "case_number", "court_name", "court_type", "region", "instance", "category", "plaintiffs",
    "defendants", "third_parties", "judge", "hearing_at", "location", "subject", "claim_amount", "url",
]

# Jadval sarlavhalarini maydonlarga moslash uchun kalit so'zlar (lotin, kirill, rus)
HEADER_KEYWORDS = {
    "case_number": ["ish raqami", "ish raqam", "ish n", "иш раками", "номер дела", "дело"],
    "court_name": ["sud nomi", "sud", "суд номи", "суд"],
    "category": ["ish turi", "ish turkumi", "toifa", "turkum", "иш тури", "категория", "вид дела"],
    "plaintiffs": ["davogar", "arizachi", "даъвогар", "истец", "заявитель"],
    "defendants": ["javobgar", "жавобгар", "ответчик"],
    "third_parties": ["uchinchi shaxs", "ishtirokchi", "учинчи", "третье лицо", "третьи лица"],
    "judge": ["sudya", "раислик", "судья"],
    "hearing_at": ["majlis vaqti", "vaqti", "sana", "мажлис", "вакт", "сана", "время", "дата"],
    "location": ["zal", "xona", "зал"],
    "subject": ["predmet", "mazmuni", "талаб", "предмет", "суть"],
    "instance": ["instansiya", "инстанция"],
    "claim_amount": ["summa", "сумма", "summasi"],
}

DATE_FORMATS = [
    "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y %H:%M",
    "%d.%m.%Y %H:%M:%S", "%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y %H:%M", "%d/%m/%Y",
]


class AdapterError(Exception):
    pass


def parse_datetime_value(value, default_date=None):
    """Turli ko'rinishdagi sana/vaqtni timezone bilan datetime'ga aylantiradi."""
    if not value:
        return None
    if isinstance(value, datetime):
        dt = value
    else:
        text = str(value).strip().replace("\xa0", " ")
        text = re.sub(r"\s+", " ", text)
        if text.endswith("Z"):
            text = text[:-1]
        text = re.sub(r"[+-]\d{2}:\d{2}$", "", text)
        text = text.split(".")[0] if re.match(r"^\d{4}-\d{2}-\d{2}T[\d:]+\.\d+$", text) else text
        dt = None
        for fmt in DATE_FORMATS:
            try:
                dt = datetime.strptime(text, fmt)
                break
            except ValueError:
                continue
        if dt is None:
            only_time = re.match(r"^(\d{1,2})[:.](\d{2})$", text)
            if only_time and default_date:
                dt = datetime.combine(default_date, datetime.min.time()).replace(
                    hour=int(only_time.group(1)), minute=int(only_time.group(2))
                )
            else:
                return None
    if timezone.is_naive(dt):
        dt = timezone.make_aware(dt)
    return dt


def get_path(data, path):
    """"a.b.0.c" ko'rinishidagi yo'l bo'yicha qiymat oladi."""
    current = data
    for part in str(path).split("."):
        if part == "":
            continue
        if isinstance(current, list):
            try:
                current = current[int(part)]
            except (ValueError, IndexError):
                return None
        elif isinstance(current, dict):
            current = current.get(part)
        else:
            return None
        if current is None:
            return None
    return current


def _join(value):
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        parts = []
        for v in value:
            if isinstance(v, dict):
                v = v.get("name") or v.get("nomi") or v.get("title") or json.dumps(v, ensure_ascii=False)
            parts.append(str(v).strip())
        return "; ".join(p for p in parts if p)
    return str(value).strip()


def clean_record(raw):
    record = {field: _join(raw.get(field)) for field in FIELDS}
    return record


class BaseAdapter:
    def __init__(self, source):
        self.source = source
        self.config = source.config or {}
        self.session = requests.Session()
        self.session.headers["User-Agent"] = self.config.get("user_agent", "SudMonitoringi/1.0 (Ijtimoiy himoya milliy agentligi)")

    def fetch(self):
        raise NotImplementedError

    def http_get(self, url, **kwargs):
        try:
            response = self.session.get(url, timeout=settings.INTEGRATION_HTTP_TIMEOUT, **kwargs)
            response.raise_for_status()
        except requests.RequestException as exc:
            raise AdapterError(f"Manbaga ulanib bo‘lmadi ({url}): {exc}") from exc
        return response

    def date_range(self):
        today = timezone.localdate()
        back = int(self.config.get("days_back", 0))
        forward = int(self.config.get("days_forward", 14))
        return [today + timedelta(days=d) for d in range(-back, forward + 1)]


class JsonApiAdapter(BaseAdapter):
    """Rasmiy API / veb-servis (JSON) adapteri.

    config namunasi:
    {
      "params": {"lang": "uz"},
      "date_from_param": "date_from", "date_to_param": "date_to", "days_back": 1, "days_forward": 30,
      "search_param": "party",            # har bir tashkilot nomi bo'yicha alohida so'rov (ixtiyoriy)
      "items_path": "data.items",
      "page_param": "page", "max_pages": 20,
      "auth_header": "Authorization", "auth_prefix": "Bearer ",
      "field_map": {"case_number": "caseNumber", "court_name": "court.name", "plaintiffs": "plaintiffs",
                    "defendants": "defendants", "hearing_at": "hearingDate", "judge": "judge.fullName"}
    }
    """

    def fetch(self):
        if not self.source.base_url:
            raise AdapterError("API manzili ko‘rsatilmagan.")
        headers = {"Accept": "application/json"}
        if self.source.auth_token:
            headers[self.config.get("auth_header", "Authorization")] = (
                self.config.get("auth_prefix", "Bearer ") + self.source.auth_token
            )
        base_params = dict(self.config.get("params", {}))
        dates = self.date_range()
        if self.config.get("date_from_param"):
            base_params[self.config["date_from_param"]] = dates[0].isoformat()
        if self.config.get("date_to_param"):
            base_params[self.config["date_to_param"]] = dates[-1].isoformat()

        queries = [base_params]
        if self.config.get("search_param"):
            from core.models import Organization

            names = set()
            for org in Organization.objects.filter(is_active=True):
                names.update(n for n in [org.short_name or org.full_name] if n)
            queries = [{**base_params, self.config["search_param"]: name} for name in sorted(names)]

        field_map = self.config.get("field_map", {})
        page_param = self.config.get("page_param")
        max_pages = int(self.config.get("max_pages", 1 if not page_param else 20))
        for params in queries:
            for page in range(1, max_pages + 1):
                if page_param:
                    params = {**params, page_param: page}
                data = self.http_get(self.source.base_url, params=params, headers=headers).json()
                items = get_path(data, self.config.get("items_path", "")) if self.config.get("items_path") else data
                if not isinstance(items, list) or not items:
                    break
                for item in items:
                    raw = {field: get_path(item, path) for field, path in field_map.items()}
                    for field in FIELDS:
                        if field not in field_map and field in item:
                            raw[field] = item[field]
                    yield clean_record(raw)
                if not page_param:
                    break


class _TableParser(HTMLParser):
    """Sahifadagi barcha <table> jadvallarini qatorlar ro'yxatiga aylantiradi."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables, self._rows, self._row, self._cell = [], None, None, None
        self._depth = 0
        self._links = []

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self._depth += 1
            if self._depth == 1:
                self._rows = []
        elif tag == "tr" and self._depth == 1:
            self._row = []
        elif tag in ("td", "th") and self._row is not None and self._depth == 1:
            self._cell = []
            self._links = []
        elif tag == "a" and self._cell is not None:
            href = dict(attrs).get("href")
            if href:
                self._links.append(href)
        elif tag == "br" and self._cell is not None:
            self._cell.append("; ")

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._cell is not None and self._row is not None:
            self._row.append({"text": re.sub(r"\s+", " ", "".join(self._cell)).strip(" ;"), "links": self._links})
            self._cell = None
        elif tag == "tr" and self._row is not None and self._depth == 1:
            if self._row:
                self._rows.append(self._row)
            self._row = None
        elif tag == "table":
            if self._depth == 1 and self._rows:
                self.tables.append(self._rows)
            self._depth = max(0, self._depth - 1)

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


def map_headers(headers, overrides=None):
    """Jadval sarlavhalarini yagona maydonlarga moslashtiradi."""
    mapping = {}
    overrides = {normalize_name(k): v for k, v in (overrides or {}).items()}
    for index, header in enumerate(headers):
        key = normalize_name(header)
        if key in overrides:
            mapping[index] = overrides[key]
            continue
        best, best_len = None, 0
        for field, words in HEADER_KEYWORDS.items():
            if field in mapping.values():
                continue
            for word in words:
                w = normalize_name(word)
                if w and w in key and len(w) > best_len:
                    best, best_len = field, len(w)
        if best:
            mapping[index] = best
    return mapping


def parse_html_tables(html, base_url="", header_overrides=None, default_date=None, court_name=""):
    parser = _TableParser()
    parser.feed(html)
    records = []
    for rows in parser.tables:
        header_row = rows[0]
        mapping = map_headers([c["text"] for c in header_row], header_overrides)
        if "case_number" not in mapping.values():
            continue
        for row in rows[1:]:
            raw = {}
            for index, field in mapping.items():
                if index < len(row):
                    raw[field] = row[index]["text"]
                    if not raw.get("url") and row[index]["links"]:
                        raw["url"] = urljoin(base_url, row[index]["links"][0])
            if not raw.get("case_number"):
                continue
            if raw.get("hearing_at"):
                dt = parse_datetime_value(raw["hearing_at"], default_date)
                raw["hearing_at"] = dt.strftime("%Y-%m-%dT%H:%M") if dt else ""
            if court_name and not raw.get("court_name"):
                raw["court_name"] = court_name
            records.append(clean_record(raw))
    return records


class HtmlTableAdapter(BaseAdapter):
    """jadval2.sud.uz sahifalaridagi sud majlislari jadvalini o'qiydi.

    config namunasi:
    {
      "pages": [
        {"url": "https://jadval2.sud.uz/...?date={date}", "court_name": "...", "court_type": "...", "region": "..."}
      ],
      "date_format": "%d.%m.%Y", "days_back": 0, "days_forward": 14,
      "header_map": {"Ish raqami": "case_number"}
    }
    `pages` berilmasa `base_url` ning o'zi o'qiladi. URL ichidagi {date} o'rniga sana qo'yiladi.
    """

    def fetch(self):
        pages = self.config.get("pages") or [{"url": self.source.base_url}]
        date_format = self.config.get("date_format", "%d.%m.%Y")
        for page in pages:
            if isinstance(page, str):
                page = {"url": page}
            url_template = page.get("url") or self.source.base_url
            if not url_template:
                continue
            dates = self.date_range() if "{date}" in url_template else [None]
            for day in dates:
                url = url_template.replace("{date}", day.strftime(date_format)) if day else url_template
                html = self.http_get(url).text
                for record in parse_html_tables(
                    html, url, self.config.get("header_map"), default_date=day, court_name=page.get("court_name", "")
                ):
                    for extra in ("court_type", "region", "instance"):
                        if page.get(extra) and not record.get(extra):
                            record[extra] = page[extra]
                    yield record


class DemoAdapter(BaseAdapter):
    """Sinov uchun manba: tizimdagi tashkilotlar ishtirokidagi soxta sud ishlarini qaytaradi.

    Har chaqiruvda ba'zi majlis sanalari o'zgaradi — o'zgarishlarni aniqlash mexanizmini ko'rsatish uchun.
    """

    PEOPLE = [
        "Karimov Anvar Tohirovich", "Rahimova Dilnoza Sobirovna", "Yusupov Jamshid Olimovich", "Tursunova Malika Baxtiyorovna",
        "Ergashev Bobur Nurmatovich", "Qodirova Nigora Akmalovna", "Sodiqov Farrux Ilhomovich", "Abdullayeva Zarina Rustamovna",
    ]
    OTHER = ["“Barakali hosil” MChJ", "“Sharq qurilish” AJ", "Toshkent shahar hokimligi", "Moliya vazirligi"]

    def fetch(self):
        from core.models import Organization

        orgs = list(Organization.objects.filter(is_active=True).exclude(org_type__code="agentlik")[:60])
        count = int(self.config.get("count", 12))
        rnd = random.Random(int(self.config.get("seed", 2026)))
        churn = random.Random()
        today = timezone.localdate()
        for i in range(count):
            person = rnd.choice(self.PEOPLE)
            if orgs and rnd.random() < 0.85:
                org = rnd.choice(orgs)
                org_name = org.full_name if rnd.random() < 0.6 else (org.alt_names.splitlines() or [org.full_name])[0]
            else:
                org_name = rnd.choice(self.OTHER)
            role = rnd.choice(["plaintiff", "defendant", "defendant", "third"])
            plaintiffs, defendants, third = person, org_name, ""
            if role == "plaintiff":
                plaintiffs, defendants = org_name, person
            elif role == "third":
                plaintiffs, defendants, third = person, rnd.choice(self.OTHER), org_name
            court_kind = rnd.choice(["Fuqarolik ishlari bo‘yicha", "Iqtisodiy", "Ma’muriy"])
            region = rnd.choice(["Toshkent shahar", "Samarqand viloyat", "Farg‘ona viloyat", "Andijon viloyat"])
            day = today + timedelta(days=rnd.randint(-3, 25) + (churn.randint(0, 3) if i % 4 == 0 else 0))
            yield clean_record({
                "external_key": f"demo-{i}",
                "case_number": f"2-{1000 + rnd.randint(1, 99)}-2601/{rnd.randint(100, 9999)}",
                "court_name": f"{court_kind} {region} sudi",
                "court_type": court_kind,
                "region": region,
                "instance": "Birinchi instansiya",
                "category": rnd.choice(["Mehnat nizolari", "Ijtimoiy nafaqa va moddiy yordam bo‘yicha nizolar", "Pul mablag‘larini undirish"]),
                "plaintiffs": plaintiffs,
                "defendants": defendants,
                "third_parties": third,
                "judge": rnd.choice(["A. Xolmatov", "S. Nazarova", "D. Ismoilov", "M. Rasulova"]),
                "hearing_at": f"{day.isoformat()}T{rnd.choice(['09:00', '10:30', '14:00', '15:30'])}",
                "location": f"{rnd.randint(1, 12)}-zal",
                "subject": "Sinov yozuvi (demo manba)",
            })


def parse_uploaded_file(uploaded):
    """Fayldan import: JSON (ro'yxat), CSV yoki Excel. Ustun nomlari maydon nomlari yoki sarlavhalar bo'lishi mumkin."""
    name = uploaded.name.lower()
    content = uploaded.read()
    rows = []
    if name.endswith(".json"):
        data = json.loads(content.decode("utf-8-sig"))
        if isinstance(data, dict):
            data = data.get("items") or data.get("data") or []
        rows = [r for r in data if isinstance(r, dict)]
    elif name.endswith(".csv"):
        text = content.decode("utf-8-sig")
        dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t") if text.strip() else csv.excel
        rows = list(csv.DictReader(io.StringIO(text), dialect=dialect))
    elif name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        values = list(ws.iter_rows(values_only=True))
        if values:
            headers = [str(h or "").strip() for h in values[0]]
            for row in values[1:]:
                if any(v not in (None, "") for v in row):
                    rows.append({headers[i]: row[i] for i in range(min(len(headers), len(row)))})
    else:
        raise AdapterError("Qo‘llab-quvvatlanadigan formatlar: JSON, CSV, XLSX.")

    records = []
    for row in rows:
        headers = list(row.keys())
        if all(h in FIELDS for h in headers if h):
            raw = dict(row)
        else:
            mapping = map_headers(headers)
            raw = {mapping[i]: row[h] for i, h in enumerate(headers) if i in mapping}
            for h in headers:
                if h in FIELDS:
                    raw[h] = row[h]
        for key, value in list(raw.items()):
            if isinstance(value, (datetime, date)):
                raw[key] = value.strftime("%Y-%m-%dT%H:%M") if isinstance(value, datetime) else value.isoformat()
        if raw.get("hearing_at"):
            dt = parse_datetime_value(raw["hearing_at"])
            raw["hearing_at"] = dt.strftime("%Y-%m-%dT%H:%M") if dt else ""
        record = clean_record(raw)
        if record["case_number"]:
            records.append(record)
    return records


ADAPTERS = {
    "json_api": JsonApiAdapter,
    "html": HtmlTableAdapter,
    "demo": DemoAdapter,
}


def get_adapter(source):
    cls = ADAPTERS.get(source.adapter)
    if cls is None:
        raise AdapterError(f"Noma’lum ulanish usuli: {source.adapter}")
    return cls(source)
