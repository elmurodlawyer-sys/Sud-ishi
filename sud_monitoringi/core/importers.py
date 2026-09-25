"""Tashkilotlar ma'lumotnomasini Excel fayldan import qilish.

Ustunlar (birinchi qator sarlavha):
    full_name | short_name | stir | org_type | region | district | parent | previous_names | alt_names | is_active

* org_type va region — klassifikator kodi yoki nomi;
* parent — yuqori turuvchi tashkilot STIR yoki to'liq nomi;
* previous_names / alt_names — bir nechta qiymat ";" bilan ajratiladi.
STIR bo'yicha (bo'lmasa to'liq nom bo'yicha) mavjud tashkilot yangilanadi.
"""
import io

from django.db import transaction
from openpyxl import load_workbook

from .models import Classifier, ClassifierKind, Organization
from .normalize import normalize_name

COLUMNS = ["full_name", "short_name", "stir", "org_type", "region", "district", "parent", "previous_names", "alt_names", "is_active"]
HEADER_ALIASES = {
    "to'liq nomi": "full_name", "toliq nomi": "full_name", "nomi": "full_name", "qisqa nomi": "short_name",
    "qisqartirilgan nomi": "short_name", "stir": "stir", "inn": "stir", "turi": "org_type", "tashkilot turi": "org_type",
    "hudud": "region", "viloyat": "region", "tuman": "district", "tuman (shahar)": "district",
    "yuqori tashkilot": "parent", "yuqori turuvchi tashkilot": "parent", "oldingi nomlari": "previous_names",
    "muqobil nomlari": "alt_names", "faol": "is_active",
}


def _classifier(kind, value):
    if not value:
        return None
    value = str(value).strip()
    found = Classifier.objects.filter(kind=kind, code=value).first()
    if found:
        return found
    key = normalize_name(value)
    for c in Classifier.objects.filter(kind=kind):
        if normalize_name(c.name) == key:
            return c
    return None


def _lines(value):
    return "\n".join(p.strip() for p in str(value or "").replace("\n", ";").split(";") if p.strip())


def import_organizations(fileobj):
    wb = load_workbook(io.BytesIO(fileobj.read()), read_only=True, data_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    if not rows:
        return {"created": 0, "updated": 0, "errors": ["Fayl bo‘sh."]}
    headers = []
    for h in rows[0]:
        h = str(h or "").strip()
        headers.append(h if h in COLUMNS else HEADER_ALIASES.get(h.lower().replace("‘", "'").replace("’", "'"), ""))
    created = updated = 0
    errors = []
    pending_parents = []
    with transaction.atomic():
        for line_no, row in enumerate(rows[1:], start=2):
            data = {headers[i]: row[i] for i in range(min(len(headers), len(row))) if headers[i]}
            name = str(data.get("full_name") or "").strip()
            if not name:
                continue
            stir = str(data.get("stir") or "").strip().split(".")[0]
            org_type = _classifier(ClassifierKind.ORG_TYPE, data.get("org_type"))
            if org_type is None:
                errors.append(f"{line_no}-qator: tashkilot turi aniqlanmadi ({data.get('org_type')}).")
                continue
            org = Organization.objects.filter(stir=stir).first() if stir else None
            if org is None:
                org = Organization.objects.filter(full_name__iexact=name).first()
            is_new = org is None
            org = org or Organization()
            org.full_name = name
            org.short_name = str(data.get("short_name") or "").strip()
            org.stir = stir if stir.isdigit() and len(stir) == 9 else ""
            if stir and not org.stir:
                errors.append(f"{line_no}-qator: STIR noto‘g‘ri ({stir}), bo‘sh qoldirildi.")
            org.org_type = org_type
            org.region = _classifier(ClassifierKind.REGION, data.get("region"))
            org.district = str(data.get("district") or "").strip()
            org.previous_names = _lines(data.get("previous_names"))
            org.alt_names = _lines(data.get("alt_names"))
            active = data.get("is_active")
            org.is_active = True if active in (None, "") else str(active).strip().lower() in ("1", "true", "ha", "faol", "yes")
            org.save()
            if data.get("parent"):
                pending_parents.append((org, str(data["parent"]).strip(), line_no))
            created += int(is_new)
            updated += int(not is_new)
        for org, parent_ref, line_no in pending_parents:
            parent = Organization.objects.filter(stir=parent_ref).first() or Organization.objects.filter(full_name__iexact=parent_ref).first()
            if parent is None or parent.pk in org.descendant_ids():
                errors.append(f"{line_no}-qator: yuqori tashkilot topilmadi yoki noto‘g‘ri ({parent_ref}).")
                continue
            org.parent = parent
            org.save(update_fields=["parent"])
    return {"created": created, "updated": updated, "errors": errors}
